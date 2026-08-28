"""Server-rendered customer project pages."""
from __future__ import annotations

import json
import uuid
from datetime import datetime, timedelta, timezone
from io import BytesIO

from flask import abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, select

from extensions import db
from models import User
from shared.models import AuditEvent, Organization, OrganizationMembership
from shared.notifications import cancel_pending_notifications

from . import customer_projects_bp
from .models import (
    Customer,
    CustomerContact,
    CustomerProject,
    MaterialCompetitor,
    ProjectActivity,
    ProjectMaterial,
    ProjectMember,
    ProjectReminderOverride,
    ProjectStageEvent,
    ProjectStatusCatalog,
)
from .permissions import (
    ADMIN_ROLES,
    apply_project_scope,
    can_view_project,
    current_membership,
    can_edit_prices,
    can_write,
    module_required,
    require_manager,
    require_price_edit,
    require_write,
)
from .services.projects import (
    DomainError,
    VersionConflict,
    add_activity,
    add_audit,
    add_competitor,
    add_contact,
    add_material,
    add_project_member,
    create_customer,
    create_project,
    local_day_bounds,
    restore_project,
    soft_delete_competitor,
    soft_delete_material,
    soft_delete_project,
    transition_stage,
    update_competitor,
    update_project,
    update_customer_grade,
    update_material_commercial,
)


def _membership() -> OrganizationMembership:
    membership = current_membership()
    if membership is None:
        abort(403)
    return membership


def _project_or_404(project_id: str, membership: OrganizationMembership) -> CustomerProject:
    project = db.session.scalar(
        select(CustomerProject).where(
            CustomerProject.id == project_id,
            CustomerProject.organization_id == membership.organization_id,
            CustomerProject.deleted_at.is_(None),
        )
    )
    if project is None or not can_view_project(membership, project):
        abort(404)
    return project


def _handle_domain_error(exc: DomainError, endpoint: str, **values: str):
    db.session.rollback()
    if isinstance(exc, VersionConflict) and values.get("project_id"):
        project = db.session.get(CustomerProject, values["project_id"])
        return render_template(
            "customer_projects/conflict.html",
            project=project,
            submitted=request.form.to_dict(),
            current_version=exc.current_version,
        ), 409
    flash(exc.message, "warning" if isinstance(exc, VersionConflict) else "danger")
    return redirect(url_for(endpoint, **values))


@customer_projects_bp.get("/")
@login_required
@module_required
def dashboard():
    membership = _membership()
    now = datetime.now(timezone.utc)
    organization = db.session.get(Organization, membership.organization_id)
    day_start, day_end = local_day_bounds(
        now, organization.timezone if organization else "Asia/Shanghai"
    )
    base = apply_project_scope(
        select(CustomerProject).where(CustomerProject.deleted_at.is_(None)), membership
    )
    projects = list(
        db.session.scalars(base.order_by(CustomerProject.next_follow_up_at.asc()).limit(12))
    )
    active_stages = {"evaluation", "initiated", "sampling", "pilot_batch", "trial_production", "design_win"}
    def scoped_count(*conditions) -> int:
        statement = apply_project_scope(
            select(func.count()).select_from(CustomerProject).where(
                CustomerProject.deleted_at.is_(None), *conditions
            ),
            membership,
        )
        return int(db.session.scalar(statement) or 0)

    stale_rules = list(
        db.session.scalars(
            select(ProjectStatusCatalog).where(
                ProjectStatusCatalog.organization_id == membership.organization_id,
                ProjectStatusCatalog.stale_after_days.is_not(None),
                ProjectStatusCatalog.is_active.is_(True),
            )
        )
    )
    stale_condition = or_(
        *[
            (CustomerProject.stage_code == rule.code)
            & (CustomerProject.last_meaningful_update_at < now - timedelta(days=rule.stale_after_days))
            for rule in stale_rules
        ]
    ) if stale_rules else (CustomerProject.id.is_(None))
    counts = {
        "overdue": scoped_count(
            CustomerProject.stage_code.in_(active_stages),
            CustomerProject.next_follow_up_at < day_start,
        ),
        "today": scoped_count(
            CustomerProject.stage_code.in_(active_stages),
            CustomerProject.next_follow_up_at >= day_start,
            CustomerProject.next_follow_up_at < day_end,
        ),
        "upcoming": scoped_count(
            CustomerProject.stage_code.in_(active_stages),
            CustomerProject.next_follow_up_at >= day_end,
            CustomerProject.next_follow_up_at < day_end + timedelta(days=7),
        ),
        "stale": scoped_count(stale_condition),
    }
    customer_names = _customer_name_map(projects)
    user_names = _user_name_map({p.primary_sales_user_id for p in projects})
    stages = _stage_map(membership.organization_id)
    return render_template(
        "customer_projects/dashboard.html",
        projects=projects,
        counts=counts,
        customer_names=customer_names,
        user_names=user_names,
        stages=stages,
    )


def _aware(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _customer_name_map(projects: list[CustomerProject]) -> dict[str, str]:
    ids = {p.customer_id for p in projects}
    if not ids:
        return {}
    return {row.id: row.name for row in db.session.scalars(select(Customer).where(Customer.id.in_(ids)))}


def _user_name_map(ids: set[int]) -> dict[int, str]:
    if not ids:
        return {}
    return {row.id: row.display_name for row in db.session.scalars(select(User).where(User.id.in_(ids)))}


def _stage_map(organization_id: str) -> dict[str, str]:
    return {
        row.code: row.display_name
        for row in db.session.scalars(
            select(ProjectStatusCatalog)
            .where(ProjectStatusCatalog.organization_id == organization_id)
            .order_by(ProjectStatusCatalog.sort_order)
        )
    }


def _filtered_project_statement(
    membership: OrganizationMembership, q: str, stage: str
):
    statement = apply_project_scope(
        select(CustomerProject).where(CustomerProject.deleted_at.is_(None)), membership
    )
    if q:
        like = f"%{q[:100]}%"
        customer_ids = select(Customer.id).where(
            Customer.organization_id == membership.organization_id,
            Customer.name.ilike(like),
        )
        material_project_ids = select(ProjectMaterial.project_id).where(
            ProjectMaterial.organization_id == membership.organization_id,
            or_(
                ProjectMaterial.promoted_mpn.ilike(like),
                ProjectMaterial.promoted_brand.ilike(like),
            ),
        )
        competitor_project_ids = (
            select(ProjectMaterial.project_id)
            .join(MaterialCompetitor, MaterialCompetitor.project_material_id == ProjectMaterial.id)
            .where(
                ProjectMaterial.organization_id == membership.organization_id,
                MaterialCompetitor.organization_id == membership.organization_id,
                or_(MaterialCompetitor.mpn.ilike(like), MaterialCompetitor.brand.ilike(like)),
            )
        )
        statement = statement.where(
            or_(
                CustomerProject.name.ilike(like),
                CustomerProject.product_name.ilike(like),
                CustomerProject.project_code.ilike(like),
                CustomerProject.customer_id.in_(customer_ids),
                CustomerProject.id.in_(material_project_ids),
                CustomerProject.id.in_(competitor_project_ids),
            )
        )
    if stage:
        statement = statement.where(CustomerProject.stage_code == stage)
    return statement


@customer_projects_bp.get("/projects")
@login_required
@module_required
def projects():
    membership = _membership()
    q = request.args.get("q", "").strip()
    stage = request.args.get("stage", "").strip()
    statement = _filtered_project_statement(membership, q, stage)
    rows = list(db.session.scalars(statement.order_by(CustomerProject.updated_at.desc()).limit(100)))
    return render_template(
        "customer_projects/projects.html",
        projects=rows,
        customer_names=_customer_name_map(rows),
        user_names=_user_name_map({p.primary_sales_user_id for p in rows}),
        stages=_stage_map(membership.organization_id),
        q=q,
        selected_stage=stage,
    )


@customer_projects_bp.get("/projects/export.xlsx")
@login_required
@module_required
def projects_export():
    membership = _membership()
    q = request.args.get("q", "").strip()
    stage = request.args.get("stage", "").strip()
    projects = list(
        db.session.scalars(
            _filtered_project_statement(membership, q, stage).order_by(
                CustomerProject.updated_at.desc(), CustomerProject.id
            )
        )
    )
    customer_names = _customer_name_map(projects)
    user_names = _user_name_map({project.primary_sales_user_id for project in projects})
    project_ids = [project.id for project in projects]
    materials = (
        list(
            db.session.scalars(
                select(ProjectMaterial)
                .where(
                    ProjectMaterial.organization_id == membership.organization_id,
                    ProjectMaterial.project_id.in_(project_ids),
                    ProjectMaterial.deleted_at.is_(None),
                )
                .order_by(ProjectMaterial.project_id, ProjectMaterial.is_primary.desc())
            )
        )
        if project_ids
        else []
    )
    materials_by_project: dict[str, list[ProjectMaterial]] = {project_id: [] for project_id in project_ids}
    for material in materials:
        materials_by_project.setdefault(material.project_id, []).append(material)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户项目台账"
    headers = [
        "项目编号", "客户", "客户评级", "项目名称", "产品名称", "项目年用量", "阶段",
        "主业务", "下一步", "下次跟进", "推广品牌", "推广型号", "应用位置", "单机数量",
        "录入单价", "录入币别", "美元单价", "含税人民币单价", "USD/CNY汇率", "价格更新时间",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D6EFD")
    customer_by_id = {
        customer.id: customer
        for customer in db.session.scalars(
            select(Customer).where(Customer.id.in_({project.customer_id for project in projects}))
        )
    } if projects else {}
    stage_names = _stage_map(membership.organization_id)
    for project in projects:
        project_materials = materials_by_project.get(project.id) or [None]
        customer = customer_by_id.get(project.customer_id)
        for material in project_materials:
            sheet.append([
                project.project_code,
                customer_names.get(project.customer_id, ""),
                customer.grade if customer else "",
                project.name,
                project.product_name or "",
                float(project.annual_usage) if project.annual_usage is not None else None,
                stage_names.get(project.stage_code, project.stage_code),
                user_names.get(project.primary_sales_user_id, ""),
                project.next_action,
                _aware(project.next_follow_up_at).strftime("%Y-%m-%d %H:%M"),
                material.promoted_brand if material else "",
                (material.promoted_mpn or "待确认") if material else "",
                (material.application_position or "") if material else "",
                float(material.machine_quantity) if material and material.machine_quantity is not None else None,
                float(material.target_price) if material and material.target_price is not None else None,
                (material.currency or "") if material else "",
                float(material.unit_price_usd) if material and material.unit_price_usd is not None else None,
                float(material.unit_price_cny_tax_included) if material and material.unit_price_cny_tax_included is not None else None,
                float(material.fx_rate_usd_cny) if material and material.fx_rate_usd_cny is not None else None,
                _aware(material.price_updated_at).strftime("%Y-%m-%d %H:%M") if material and material.price_updated_at else "",
            ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)
        sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    add_audit(
        membership.organization_id,
        "customer_project_export",
        membership.organization_id,
        "exported",
        membership.user_id,
        {"project_count": len(projects), "material_count": len(materials), "filtered": bool(q or stage)},
    )
    db.session.commit()
    return send_file(
        output,
        as_attachment=True,
        download_name=f"客户项目台账-{datetime.now().strftime('%Y%m%d')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@customer_projects_bp.route("/customers", methods=["GET", "POST"])
@login_required
@module_required
def customers():
    membership = _membership()
    require_write(membership)
    if request.method == "POST":
        try:
            customer = create_customer(request.form.to_dict(), membership)
            db.session.commit()
            flash(f"客户“{customer.name}”已创建。", "success")
            return redirect(url_for("customer_projects.customers"))
        except DomainError as exc:
            return _handle_domain_error(exc, "customer_projects.customers")
    statement = select(Customer).where(
        Customer.organization_id == membership.organization_id,
        Customer.deleted_at.is_(None),
    )
    if not membership.roles.intersection(ADMIN_ROLES):
        statement = statement.where(Customer.primary_owner_user_id == membership.user_id)
    rows = list(db.session.scalars(statement.order_by(Customer.name).limit(100)))
    customer_ids = [row.id for row in rows]
    contacts = list(
        db.session.scalars(
            select(CustomerContact)
            .where(
                CustomerContact.customer_id.in_(customer_ids),
                CustomerContact.deleted_at.is_(None),
            )
            .order_by(CustomerContact.is_primary.desc(), CustomerContact.name)
        )
    ) if customer_ids else []
    contacts_by_customer: dict[str, list[CustomerContact]] = {customer_id: [] for customer_id in customer_ids}
    for contact in contacts:
        contacts_by_customer.setdefault(contact.customer_id, []).append(contact)
    return render_template("customer_projects/customers.html", customers=rows, contacts_by_customer=contacts_by_customer)


@customer_projects_bp.post("/customers/<string:customer_id>/grade")
@login_required
@module_required
def customer_update_grade(customer_id: str):
    membership = _membership()
    require_write(membership)
    customer = db.session.scalar(
        select(Customer).where(
            Customer.id == customer_id,
            Customer.organization_id == membership.organization_id,
            Customer.deleted_at.is_(None),
        )
    )
    if customer is None or (
        not membership.roles.intersection(ADMIN_ROLES)
        and customer.primary_owner_user_id != membership.user_id
    ):
        abort(404)
    try:
        update_customer_grade(customer, request.form.get("grade"), membership)
        db.session.commit()
        flash("客户评级已更新。", "success")
        return redirect(url_for("customer_projects.customers"))
    except DomainError as exc:
        return _handle_domain_error(exc, "customer_projects.customers")


@customer_projects_bp.post("/customers/<string:customer_id>/contacts")
@login_required
@module_required
def customer_add_contact(customer_id: str):
    membership = _membership()
    require_write(membership)
    customer = db.session.scalar(
        select(Customer).where(
            Customer.id == customer_id,
            Customer.organization_id == membership.organization_id,
            Customer.deleted_at.is_(None),
        )
    )
    if customer is None or (
        not membership.roles.intersection(ADMIN_ROLES)
        and customer.primary_owner_user_id != membership.user_id
    ):
        abort(404)
    try:
        data = request.form.to_dict()
        data["is_primary"] = request.form.get("is_primary") == "on"
        add_contact(customer, data, membership)
        db.session.commit()
        flash("联系人已添加。", "success")
        return redirect(url_for("customer_projects.customers"))
    except DomainError as exc:
        return _handle_domain_error(exc, "customer_projects.customers")


@customer_projects_bp.route("/projects/new", methods=["GET", "POST"])
@login_required
@module_required
def project_new():
    membership = _membership()
    require_write(membership)
    if request.method == "POST":
        try:
            project = create_project(
                request.form.to_dict(), membership, request.form.get("idempotency_key", "")
            )
            db.session.commit()
            flash(f"项目 {project.project_code} 已创建。", "success")
            return redirect(url_for("customer_projects.project_detail", project_id=project.id))
        except DomainError as exc:
            return _handle_domain_error(exc, "customer_projects.project_new")
    customer_stmt = select(Customer).where(Customer.organization_id == membership.organization_id, Customer.deleted_at.is_(None))
    if not membership.roles.intersection(ADMIN_ROLES):
        customer_stmt = customer_stmt.where(Customer.primary_owner_user_id == membership.user_id)
    customers = list(db.session.scalars(customer_stmt.order_by(Customer.name)))
    members = list(
        db.session.execute(
            select(User, OrganizationMembership)
            .join(OrganizationMembership, OrganizationMembership.user_id == User.id)
            .where(
                OrganizationMembership.organization_id == membership.organization_id,
                OrganizationMembership.status == "active",
                User.is_active_user.is_(True),
            )
            .order_by(User.email)
        )
    )
    return render_template(
        "customer_projects/project_new.html",
        customers=customers,
        members=members,
        stages=_stage_map(membership.organization_id),
        idempotency_key=str(uuid.uuid4()),
    )


@customer_projects_bp.get("/projects/<string:project_id>")
@login_required
@module_required
def project_detail(project_id: str):
    membership = _membership()
    project = _project_or_404(project_id, membership)
    customer = db.session.get(Customer, project.customer_id)
    materials = list(db.session.scalars(select(ProjectMaterial).where(ProjectMaterial.project_id == project.id, ProjectMaterial.deleted_at.is_(None)).order_by(ProjectMaterial.is_primary.desc(), ProjectMaterial.created_at)))
    material_ids = [m.id for m in materials]
    competitors = list(db.session.scalars(select(MaterialCompetitor).where(MaterialCompetitor.project_material_id.in_(material_ids), MaterialCompetitor.deleted_at.is_(None)))) if material_ids else []
    competitors_by_material: dict[str, list[MaterialCompetitor]] = {mid: [] for mid in material_ids}
    for competitor in competitors:
        competitors_by_material.setdefault(competitor.project_material_id, []).append(competitor)
    activities = list(db.session.scalars(select(ProjectActivity).where(ProjectActivity.project_id == project.id).order_by(ProjectActivity.occurred_at.desc()).limit(100)))
    stage_events = list(db.session.scalars(select(ProjectStageEvent).where(ProjectStageEvent.project_id == project.id).order_by(ProjectStageEvent.occurred_at.desc()).limit(100)))
    timeline = [
        {"kind": "activity", "occurred_at": item.occurred_at, "item": item}
        for item in activities
    ] + [
        {"kind": "stage", "occurred_at": item.occurred_at, "item": item}
        for item in stage_events
    ]
    timeline.sort(key=lambda row: _aware(row["occurred_at"]), reverse=True)
    audits = []
    if membership.roles.intersection(ADMIN_ROLES):
        audits = list(db.session.scalars(select(AuditEvent).where(AuditEvent.organization_id == membership.organization_id, AuditEvent.object_type == "project", AuditEvent.object_id == project.id).order_by(AuditEvent.occurred_at.desc()).limit(50)))
    project_members = list(
        db.session.execute(
            select(ProjectMember, User)
            .join(User, User.id == ProjectMember.user_id)
            .where(ProjectMember.project_id == project.id, ProjectMember.left_at.is_(None))
            .order_by(ProjectMember.role_code, User.email)
        )
    )
    organization_members = []
    reminder_override = db.session.scalar(
        select(ProjectReminderOverride).where(
            ProjectReminderOverride.organization_id == membership.organization_id,
            ProjectReminderOverride.project_id == project.id,
        )
    )
    member_email_preferences = {}
    for item, _user in project_members:
        try:
            value = json.loads(item.notification_preferences_json or "{}")
        except (TypeError, ValueError):
            value = {}
        member_email_preferences[item.id] = not isinstance(value, dict) or value.get("email_enabled", True) is not False
    if membership.roles.intersection(ADMIN_ROLES):
        organization_members = list(
            db.session.execute(
                select(User, OrganizationMembership)
                .join(OrganizationMembership, OrganizationMembership.user_id == User.id)
                .where(
                    OrganizationMembership.organization_id == membership.organization_id,
                    OrganizationMembership.status == "active",
                    User.is_active_user.is_(True),
                )
                .order_by(User.email)
            )
        )
    return render_template(
        "customer_projects/project_detail.html",
        project=project,
        customer=customer,
        materials=materials,
        competitors_by_material=competitors_by_material,
        activities=activities,
        stage_events=stage_events,
        timeline=timeline,
        audits=audits,
        project_members=project_members,
        member_email_preferences=member_email_preferences,
        organization_members=organization_members,
        stages=_stage_map(membership.organization_id),
        can_manage=bool(membership.roles.intersection(ADMIN_ROLES)),
        can_write=can_write(membership),
        can_edit_price=can_edit_prices(membership),
        reminder_override=reminder_override,
        form_key=str(uuid.uuid4()),
    )



@customer_projects_bp.post("/projects/<string:project_id>/edit")
@login_required
@module_required
def project_edit(project_id: str):
    membership = _membership()
    require_write(membership)
    project = _project_or_404(project_id, membership)
    try:
        updated = update_project(
            project,
            request.form.to_dict(),
            membership,
            int(request.form.get("project_version", "0")),
        )
        db.session.commit()
        flash("项目基础信息已更新。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=updated.id))
    except (DomainError, ValueError) as exc:
        if isinstance(exc, ValueError) and not isinstance(exc, DomainError):
            exc = DomainError("VALIDATION_ERROR", "项目版本或概率区间无效。")
        return _handle_domain_error(
            exc, "customer_projects.project_detail", project_id=project.id
        )


@customer_projects_bp.post("/projects/<string:project_id>/activities")
@login_required
@module_required
def project_add_activity(project_id: str):
    membership = _membership()
    require_write(membership)
    project = _project_or_404(project_id, membership)
    try:
        data = request.form.to_dict()
        data["is_meaningful"] = request.form.get("is_meaningful") == "on"
        add_activity(project, data, membership, request.form.get("idempotency_key", ""))
        db.session.commit()
        flash("跟进记录已保存，下一步已同步更新。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except DomainError as exc:
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/projects/<string:project_id>/materials")
@login_required
@module_required
def project_add_material(project_id: str):
    membership = _membership()
    require_write(membership)
    project = _project_or_404(project_id, membership)
    try:
        data = request.form.to_dict()
        data["mpn_pending"] = request.form.get("mpn_pending") == "on"
        data["is_primary"] = request.form.get("is_primary") == "on"
        if data.get("unit_price") not in (None, ""):
            require_price_edit(membership)
        add_material(project, data, membership, request.form.get("idempotency_key", ""))
        db.session.commit()
        flash("推广物料已添加。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except DomainError as exc:
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/materials/<string:material_id>/commercial")
@login_required
@module_required
def material_update_commercial(material_id: str):
    membership = _membership()
    require_write(membership)
    material = db.session.scalar(
        select(ProjectMaterial).where(
            ProjectMaterial.id == material_id,
            ProjectMaterial.organization_id == membership.organization_id,
            ProjectMaterial.deleted_at.is_(None),
        )
    )
    if material is None:
        abort(404)
    project = _project_or_404(material.project_id, membership)
    try:
        data = request.form.to_dict()
        data["mpn_pending"] = request.form.get("mpn_pending") == "on"
        data["is_primary"] = request.form.get("is_primary") == "on"
        if data.get("unit_price") not in (None, ""):
            require_price_edit(membership)
        update_material_commercial(
            material,
            data,
            membership,
            int(request.form.get("material_version", "0")),
        )
        db.session.commit()
        flash("物料商务信息已更新。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except (DomainError, ValueError) as exc:
        if isinstance(exc, ValueError) and not isinstance(exc, DomainError):
            exc = DomainError("VALIDATION_ERROR", "物料版本无效。")
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/materials/<string:material_id>/delete")
@login_required
@module_required
def material_delete(material_id: str):
    membership = _membership()
    require_write(membership)
    material = db.session.scalar(
        select(ProjectMaterial).where(
            ProjectMaterial.id == material_id,
            ProjectMaterial.organization_id == membership.organization_id,
            ProjectMaterial.deleted_at.is_(None),
        )
    )
    if material is None:
        abort(404)
    project = _project_or_404(material.project_id, membership)
    try:
        soft_delete_material(
            material,
            request.form.get("delete_reason", ""),
            membership,
            int(request.form.get("material_version", "0")),
        )
        db.session.commit()
        flash("推广物料已移除，历史审计仍保留。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except (DomainError, ValueError) as exc:
        if isinstance(exc, ValueError) and not isinstance(exc, DomainError):
            exc = DomainError("VALIDATION_ERROR", "物料版本无效。")
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/materials/<string:material_id>/competitors")
@login_required
@module_required
def material_add_competitor(material_id: str):
    membership = _membership()
    require_write(membership)
    material = db.session.scalar(select(ProjectMaterial).where(ProjectMaterial.id == material_id, ProjectMaterial.organization_id == membership.organization_id, ProjectMaterial.deleted_at.is_(None)))
    if material is None:
        abort(404)
    project = _project_or_404(material.project_id, membership)
    try:
        data = request.form.to_dict()
        data["model_pending"] = request.form.get("model_pending") == "on"
        add_competitor(material, data, membership, request.form.get("idempotency_key", ""))
        db.session.commit()
        flash("竞争方案已添加。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except DomainError as exc:
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/competitors/<string:competitor_id>/edit")
@login_required
@module_required
def competitor_edit(competitor_id: str):
    membership = _membership()
    require_write(membership)
    competitor = db.session.scalar(
        select(MaterialCompetitor).where(
            MaterialCompetitor.id == competitor_id,
            MaterialCompetitor.organization_id == membership.organization_id,
            MaterialCompetitor.deleted_at.is_(None),
        )
    )
    if competitor is None:
        abort(404)
    material = db.session.scalar(
        select(ProjectMaterial).where(
            ProjectMaterial.id == competitor.project_material_id,
            ProjectMaterial.organization_id == membership.organization_id,
            ProjectMaterial.deleted_at.is_(None),
        )
    )
    if material is None:
        abort(404)
    project = _project_or_404(material.project_id, membership)
    try:
        data = request.form.to_dict()
        data["model_pending"] = request.form.get("model_pending") == "on"
        update_competitor(
            competitor,
            data,
            membership,
            int(request.form.get("competitor_version", "0")),
        )
        db.session.commit()
        flash("竞争方案已更新。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except (DomainError, ValueError) as exc:
        if isinstance(exc, ValueError) and not isinstance(exc, DomainError):
            exc = DomainError("VALIDATION_ERROR", "竞争方案版本无效。")
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/competitors/<string:competitor_id>/delete")
@login_required
@module_required
def competitor_delete(competitor_id: str):
    membership = _membership()
    require_write(membership)
    competitor = db.session.scalar(
        select(MaterialCompetitor).where(
            MaterialCompetitor.id == competitor_id,
            MaterialCompetitor.organization_id == membership.organization_id,
            MaterialCompetitor.deleted_at.is_(None),
        )
    )
    if competitor is None:
        abort(404)
    material = db.session.scalar(
        select(ProjectMaterial).where(
            ProjectMaterial.id == competitor.project_material_id,
            ProjectMaterial.organization_id == membership.organization_id,
            ProjectMaterial.deleted_at.is_(None),
        )
    )
    if material is None:
        abort(404)
    project = _project_or_404(material.project_id, membership)
    try:
        soft_delete_competitor(
            competitor,
            request.form.get("delete_reason", ""),
            membership,
            int(request.form.get("competitor_version", "0")),
        )
        db.session.commit()
        flash("竞争方案已移除，历史审计仍保留。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except (DomainError, ValueError) as exc:
        if isinstance(exc, ValueError) and not isinstance(exc, DomainError):
            exc = DomainError("VALIDATION_ERROR", "竞争方案版本无效。")
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/projects/<string:project_id>/stage")
@login_required
@module_required
def project_transition(project_id: str):
    membership = _membership()
    require_write(membership)
    project = _project_or_404(project_id, membership)
    target = request.form.get("to_stage_code", "")
    if target in {"mass_production", "lost"}:
        require_manager(membership)
    try:
        transition_stage(project, request.form.to_dict(), membership, request.form.get("idempotency_key", ""))
        db.session.commit()
        flash("项目阶段已更新。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except DomainError as exc:
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/projects/<string:project_id>/members")
@login_required
@module_required
def project_add_member(project_id: str):
    membership = _membership()
    require_manager(membership)
    project = _project_or_404(project_id, membership)
    try:
        add_project_member(
            project,
            int(request.form.get("user_id", "0")),
            request.form.get("role_code", ""),
            membership,
        )
        db.session.commit()
        flash("项目成员已添加。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except (DomainError, ValueError) as exc:
        if isinstance(exc, ValueError) and not isinstance(exc, DomainError):
            exc = DomainError("VALIDATION_ERROR", "请选择有效成员。")
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.post("/projects/<string:project_id>/reminder-settings")
@login_required
@module_required
def project_reminder_settings(project_id: str):
    membership = _membership()
    require_manager(membership)
    project = _project_or_404(project_id, membership)
    override = db.session.scalar(
        select(ProjectReminderOverride).where(ProjectReminderOverride.project_id == project.id)
    )
    if override is None:
        override = ProjectReminderOverride(organization_id=membership.organization_id, project_id=project.id)
        db.session.add(override)
        db.session.flush()

    def tri_state(name: str):
        value = request.form.get(name, "inherit")
        return None if value == "inherit" else value == "yes"

    override.is_enabled = request.form.get("is_enabled", "yes") == "yes"
    override.include_pm = tri_state("include_pm")
    override.include_fae = tri_state("include_fae")
    override.version += 1
    cancel_pending_notifications("customer_projects", "project", project.id)
    add_audit(project.organization_id, "project_reminder_override", override.id, "updated", membership.user_id, {"enabled": override.is_enabled, "version": override.version})
    db.session.commit()
    flash("项目提醒设置已保存，未发送的旧提醒已取消。", "success")
    return redirect(url_for("customer_projects.project_detail", project_id=project.id))


@customer_projects_bp.post("/projects/<string:project_id>/members/<string:member_id>/notifications")
@login_required
@module_required
def project_member_notifications(project_id: str, member_id: str):
    membership = _membership()
    require_manager(membership)
    project = _project_or_404(project_id, membership)
    member = db.session.scalar(
        select(ProjectMember).where(
            ProjectMember.id == member_id,
            ProjectMember.organization_id == membership.organization_id,
            ProjectMember.project_id == project.id,
            ProjectMember.left_at.is_(None),
        )
    )
    if member is None:
        abort(404)
    enabled = request.form.get("email_enabled") == "on"
    member.notification_preferences_json = json.dumps({"email_enabled": enabled}, ensure_ascii=False, sort_keys=True)
    cancel_pending_notifications("customer_projects", "project", project.id)
    add_audit(project.organization_id, "project_member", member.id, "notification_preferences_updated", membership.user_id, {"email_enabled": enabled})
    db.session.commit()
    flash("成员通知偏好已更新。", "success")
    return redirect(url_for("customer_projects.project_detail", project_id=project.id))


@customer_projects_bp.post("/projects/<string:project_id>/delete")
@login_required
@module_required
def project_delete(project_id: str):
    membership = _membership()
    require_write(membership)
    project = _project_or_404(project_id, membership)
    try:
        soft_delete_project(
            project,
            request.form.get("delete_reason", ""),
            membership,
            int(request.form.get("project_version", "0")),
        )
        db.session.commit()
        flash("项目已移入回收站，可由管理员恢复。", "success")
        return redirect(url_for("customer_projects.projects"))
    except (DomainError, ValueError) as exc:
        if isinstance(exc, ValueError) and not isinstance(exc, DomainError):
            exc = DomainError("VALIDATION_ERROR", "项目版本无效。")
        return _handle_domain_error(exc, "customer_projects.project_detail", project_id=project.id)


@customer_projects_bp.get("/trash")
@login_required
@module_required
def trash():
    membership = _membership()
    require_manager(membership)
    rows = list(
        db.session.scalars(
            select(CustomerProject)
            .where(
                CustomerProject.organization_id == membership.organization_id,
                CustomerProject.deleted_at.is_not(None),
            )
            .order_by(CustomerProject.deleted_at.desc())
            .limit(100)
        )
    )
    return render_template("customer_projects/trash.html", projects=rows)


@customer_projects_bp.post("/trash/projects/<string:project_id>/restore")
@login_required
@module_required
def project_restore(project_id: str):
    membership = _membership()
    require_manager(membership)
    project = db.session.scalar(
        select(CustomerProject).where(
            CustomerProject.id == project_id,
            CustomerProject.organization_id == membership.organization_id,
            CustomerProject.deleted_at.is_not(None),
        )
    )
    if project is None:
        abort(404)
    try:
        restore_project(project, membership)
        db.session.commit()
        flash("项目已恢复。", "success")
        return redirect(url_for("customer_projects.project_detail", project_id=project.id))
    except DomainError as exc:
        return _handle_domain_error(exc, "customer_projects.trash")
